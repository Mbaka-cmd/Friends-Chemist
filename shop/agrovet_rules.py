import csv, re
from decimal import Decimal, InvalidOperation
RULES = [
 ('Milk replacers', ['milk replacer']),
 ('Animal health', ['nillzan', 'nilzan', 'abedazole', 'piperazine', 'dawacrim', 'saline']),
 ('Crop protection', ['tingatinga', 'sevin', 'weedal', 'wedelal']),
 ('Crop nutrition', ['fastgrow']),
 ('Equipment & supplies', ['needle', 'syringe', 'glove', 'kill/trap']),
 ('Minerals & supplements', ['lick', 'block', 'maclick', 'macklick', 'maclik', 'ngombemix', 'twiga', 'udaminx', 'mineral', 'dcp', 'kraal salt', 'epsom',
                             'kaolin', 'vitamin', 'vitax', 'vital', 'beemix', 'maxilick', 'udder booster', 'diamond v']),
]
RESTRICTED = {'Animal health', 'Crop protection'}
SKU_RE = re.compile(r'^Agr-\d{3}$')

def categorise(name):
    n = name.lower()
    return next((c for c, kws in RULES if any(k in n for k in kws)), 'To be categorised')

def _clean(v):
    return '' if v is None else str(v).strip()

def _dec(v):
    s = _clean(v).replace(',', '')
    try: return Decimal(s) if s else None
    except InvalidOperation: return None

def _raw_rows(path):
    p = str(path).lower()
    if p.endswith(('.xlsx', '.xlsm')):
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb['Agrovet Product List'] if 'Agrovet Product List' in wb.sheetnames else wb.worksheets[0]
        yield from ws.iter_rows(min_row=2, values_only=True)
    else:
        with open(path, encoding='utf-8-sig', newline='') as f:
            rd = csv.reader(f, delimiter='\t' if p.endswith(('.tsv', '.txt')) else ',')
            next(rd, None)
            yield from rd

def parse_rows(path):
    rows = []
    for r in _raw_rows(path):
        r = list(r) + [None] * 10
        name = _clean(r[1])
        if not name: continue
        rows.append(dict(sku=_clean(r[0]), name=name, pack=_clean(r[2]), price=_dec(r[3]), cost=_dec(r[4]),
                         note=_clean(r[7]), image=bool(_clean(r[9]))))
    nxt = max([int(r['sku'][4:]) for r in rows if SKU_RE.match(r['sku'])] or [0]) + 1
    for r in rows:
        if not SKU_RE.match(r['sku']):
            r['assigned'] = True; r['sku'] = 'Agr-%03d' % nxt; nxt += 1
    return rows